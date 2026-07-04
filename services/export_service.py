"""Export service: TXT, MD, DOCX, PDF, EPUB, HTML, JSON.

Supports both synchronous (small) and async (large) exports.
Async exports run in a background thread with progress reported via core.cache.
"""
from __future__ import annotations

import io
import json
import logging
import os
import threading
import uuid
from pathlib import Path
from typing import Any

from sqlalchemy import select

from config import Config
from core.cache import cache
from core.db import read_session
from core.errors import NotFoundError
from models.chapter import Chapter
from models.character import Character
from models.world import WorldEntry
from services._common import current_project_id, load_json

log = logging.getLogger("asm.export")


def _safe_filename(name: str) -> str:
    import re
    safe = re.sub(r"[^\w.\- ]+", "_", name or "export")
    return safe.strip()[:180]


def export_chapter(chapter_id: str, fmt: str) -> tuple[bytes, str, str]:
    """Export a single chapter. Returns (content, mimetype, filename)."""
    fmt = fmt.lower()
    with read_session() as s:
        ch = s.get(Chapter, chapter_id)
        if not ch:
            raise NotFoundError("Chapter not found.")
        title = ch.title or "Untitled"
        content = ch.content or ""
    if fmt == "txt":
        return content.encode("utf-8"), "text/plain", f"{_safe_filename(title)}.txt"
    if fmt == "md":
        md = f"# {title}\n\n{content}\n"
        return md.encode("utf-8"), "text/markdown", f"{_safe_filename(title)}.md"
    if fmt == "html":
        html = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>{title}</title>
<style>body{{font-family:Georgia,serif;max-width:720px;margin:2rem auto;padding:1rem;line-height:1.7}}h1{{text-align:center}}</style>
</head><body><h1>{title}</h1><div>{_md_to_html(content)}</div></body></html>"""
        return html.encode("utf-8"), "text/html", f"{_safe_filename(title)}.html"
    if fmt == "docx":
        return _chapter_to_docx(title, content), \
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document", \
            f"{_safe_filename(title)}.docx"
    if fmt == "pdf":
        return _chapter_to_pdf(title, content), \
            "application/pdf", f"{_safe_filename(title)}.pdf"
    if fmt == "json":
        data = {"title": title, "content": content, "format": "asm-chapter-v1"}
        return json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8"), \
            "application/json", f"{_safe_filename(title)}.json"
    raise NotFoundError(f"Unknown format: {fmt}")


def export_manuscript(fmt: str) -> tuple[bytes, str, str]:
    """Full manuscript (concatenated chapters)."""
    fmt = fmt.lower()
    with read_session() as s:
        pid = current_project_id(s)
        chapters = list(s.scalars(
            select(Chapter).where(Chapter.project_id == pid)
            .order_by(Chapter.sort_order.asc())
        ).all())
        from models.settings import Setting
        title = Setting.get(s, "story_title", "My Story")
        author = Setting.get(s, "story_author", "Author")
    if fmt == "txt":
        parts = [f"{title}\nby {author}\n\n"]
        for ch in chapters:
            parts.append(f"\n\n{ch.title}\n\n{ch.content or ''}\n")
        return "".join(parts).encode("utf-8"), "text/plain", f"{_safe_filename(title)}.txt"
    if fmt == "md":
        parts = [f"# {title}\n\n*by {author}*\n\n"]
        for ch in chapters:
            parts.append(f"\n\n## {ch.title}\n\n{ch.content or ''}\n")
        return "".join(parts).encode("utf-8"), "text/markdown", f"{_safe_filename(title)}.md"
    if fmt == "html":
        parts = [f"""<!DOCTYPE html><html><head><meta charset="utf-8"><title>{title}</title>
<style>body{{font-family:Georgia,serif;max-width:720px;margin:2rem auto;padding:1rem;line-height:1.7}}h1{{text-align:center}}h2{{page-break-before:always;border-bottom:1px solid #ccc;padding-bottom:.3rem}}</style>
</head><body><h1>{title}</h1><p style="text-align:center"><em>by {author}</em></p>"""]
        for ch in chapters:
            parts.append(f"<h2>{ch.title}</h2><div>{_md_to_html(ch.content or '')}</div>")
        parts.append("</body></html>")
        return "".join(parts).encode("utf-8"), "text/html", f"{_safe_filename(title)}.html"
    if fmt == "docx":
        return _manuscript_to_docx(title, author, chapters), \
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document", \
            f"{_safe_filename(title)}.docx"
    if fmt == "pdf":
        return _manuscript_to_pdf(title, author, chapters), \
            "application/pdf", f"{_safe_filename(title)}.pdf"
    if fmt == "epub":
        return _manuscript_to_epub(title, author, chapters), \
            "application/epub+zip", f"{_safe_filename(title)}.epub"
    raise NotFoundError(f"Unknown format: {fmt}")


def export_character_sheet(character_id: str, fmt: str = "pdf") -> tuple[bytes, str, str]:
    with read_session() as s:
        ch = s.get(Character, character_id)
        if not ch:
            raise NotFoundError("Character not found.")
        aliases = load_json(ch.aliases, [])
        quotes = load_json(ch.philosophy_quotes, [])
        data = {
            "name": ch.name, "role": ch.role, "age": ch.age,
            "gender": ch.gender, "aliases": aliases,
            "avatar_color": ch.avatar_color,
            "physical": ch.physical, "psychology": ch.psychology,
            "background": ch.background, "philosophy": ch.philosophy,
            "philosophy_quotes": quotes, "story_role": ch.story_role,
            "voice": ch.voice, "notes": ch.notes,
        }
    if fmt == "json":
        return json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8"), \
            "application/json", f"{_safe_filename(data['name'])}.json"
    if fmt == "txt":
        parts = [f"{data['name']} ({data['role']})\n{'='*40}\n"]
        for k, v in data.items():
            if k in ("name", "avatar_color"):
                continue
            if v:
                parts.append(f"\n{k.upper()}:\n{v}\n")
        return "".join(parts).encode("utf-8"), "text/plain", f"{_safe_filename(data['name'])}.txt"
    if fmt == "pdf":
        return _character_sheet_pdf(data), "application/pdf", \
            f"{_safe_filename(data['name'])}.pdf"
    if fmt == "docx":
        return _character_sheet_docx(data), \
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document", \
            f"{_safe_filename(data['name'])}.docx"
    raise NotFoundError(f"Unknown format: {fmt}")


def export_world_bible(fmt: str = "pdf") -> tuple[bytes, str, str]:
    with read_session() as s:
        pid = current_project_id(s)
        from models.settings import Setting
        title = Setting.get(s, "story_title", "My Story")
        entries = list(s.scalars(
            select(WorldEntry).where(WorldEntry.project_id == pid)
            .order_by(WorldEntry.type.asc(), WorldEntry.name.asc())
        ).all())
    if fmt == "json":
        data = {
            "title": f"{title} — World Bible",
            "entries": [
                {"type": e.type, "name": e.name, "category": e.category,
                 "description": e.description, "content": e.content, "notes": e.notes}
                for e in entries
            ],
        }
        return json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8"), \
            "application/json", "world_bible.json"
    if fmt == "txt":
        parts = [f"{title} — World Bible\n\n"]
        by_type: dict[str, list] = {}
        for e in entries:
            by_type.setdefault(e.type, []).append(e)
        for t, group in by_type.items():
            parts.append(f"\n\n=== {t.upper()} ===\n")
            for e in group:
                parts.append(f"\n--- {e.name} ---\n{e.description or ''}\n\n{e.content or ''}\n")
        return "".join(parts).encode("utf-8"), "text/plain", "world_bible.txt"
    if fmt == "pdf":
        return _world_bible_pdf(title, entries), "application/pdf", "world_bible.pdf"
    if fmt == "docx":
        return _world_bible_docx(title, entries), \
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document", \
            "world_bible.docx"
    raise NotFoundError(f"Unknown format: {fmt}")


# ---------- PDF / DOCX / EPUB helpers ----------

def _md_to_html(text: str) -> str:
    """Minimal markdown → HTML (paragraphs + simple headings)."""
    import markdown
    return markdown.markdown(text or "", extensions=["extra", "nl2br"])


def _chapter_to_docx(title: str, content: str) -> bytes:
    from docx import Document
    doc = Document()
    doc.add_heading(title, level=1)
    for para in (content or "").split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _chapter_to_pdf(title: str, content: str) -> bytes:
    """Single chapter PDF — industry standard manuscript format.

    Uses settings: manuscript_font, manuscript_font_size, manuscript_line_spacing,
    manuscript_margins.
    """
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from reportlab.lib.enums import TA_CENTER

    # Read manuscript settings
    from core.db import read_session
    from models.settings import Setting
    with read_session() as s:
        font_name = Setting.get(s, "manuscript_font", "Courier") or "Courier"
        font_size = int(Setting.get(s, "manuscript_font_size", 12) or 12)
        line_spacing = float(Setting.get(s, "manuscript_line_spacing", 2.0) or 2.0)
        margins = float(Setting.get(s, "manuscript_margins", 1.0) or 1.0)

    leading = int(font_size * line_spacing)  # double-spaced = 24pt for 12pt font

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter,
                            leftMargin=margins*inch, rightMargin=margins*inch,
                            topMargin=margins*inch, bottomMargin=margins*inch)
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("ChapterH1", parent=styles["Heading1"],
                        alignment=TA_CENTER, fontSize=18, spaceAfter=24)
    body = ParagraphStyle("ChapterBody", parent=styles["BodyText"],
                          fontName=font_name, fontSize=font_size, leading=leading,
                          firstLineIndent=0.5*inch, spaceAfter=10)
    story = [Paragraph(title, h1), Spacer(1, 12)]
    for para in (content or "").split("\n\n"):
        if para.strip():
            safe = para.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(safe, body))
            story.append(Spacer(1, 6))
    doc.build(story)
    return buf.getvalue()


def _manuscript_to_docx(title: str, author: str, chapters: list[Chapter]) -> bytes:
    from docx import Document
    doc = Document()
    # Title page
    t = doc.add_paragraph()
    t.alignment = 1  # center
    run = t.add_run(title)
    run.bold = True
    run.font.size = __import__("docx").shared.Pt(28)
    doc.add_paragraph()
    a = doc.add_paragraph()
    a.alignment = 1
    a.add_run(f"by {author}").italic = True
    doc.add_page_break()
    for ch in chapters:
        doc.add_heading(ch.title, level=1)
        for para in (ch.content or "").split("\n\n"):
            if para.strip():
                doc.add_paragraph(para.strip())
        doc.add_page_break()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _manuscript_to_pdf(title: str, author: str, chapters: list[Chapter]) -> bytes:
    """Full manuscript PDF — industry standard manuscript format.

    Uses settings: manuscript_font, manuscript_font_size, manuscript_line_spacing,
    manuscript_margins. Default: 12pt Courier, double-spaced (24pt leading),
    1-inch margins, 0.5-inch first-line indent.
    """
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from reportlab.lib.enums import TA_CENTER

    # Read manuscript settings
    from core.db import read_session
    from models.settings import Setting
    with read_session() as s:
        font_name = Setting.get(s, "manuscript_font", "Courier") or "Courier"
        font_size = int(Setting.get(s, "manuscript_font_size", 12) or 12)
        line_spacing = float(Setting.get(s, "manuscript_line_spacing", 2.0) or 2.0)
        margins = float(Setting.get(s, "manuscript_margins", 1.0) or 1.0)

    leading = int(font_size * line_spacing)  # double-spaced = 24pt for 12pt font

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter,
                            leftMargin=margins*inch, rightMargin=margins*inch,
                            topMargin=margins*inch, bottomMargin=margins*inch)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title", parent=styles["Title"],
                                 alignment=TA_CENTER, fontSize=24, spaceAfter=18)
    author_style = ParagraphStyle("Author", parent=styles["Normal"],
                                  alignment=TA_CENTER, fontSize=14, spaceAfter=24)
    h1 = ParagraphStyle("ChH1", parent=styles["Heading1"],
                         alignment=TA_CENTER, fontSize=18, spaceAfter=24)
    body = ParagraphStyle("Body", parent=styles["BodyText"],
                          fontName=font_name, fontSize=font_size, leading=leading,
                          firstLineIndent=0.5*inch, spaceAfter=10)

    story = [
        Spacer(1, 3*inch), Paragraph(title, title_style),
        Paragraph(f"by {author}", author_style),
        Spacer(1, 2*inch),
        Paragraph(f"Approx. {sum((c.word_count or 0) for c in chapters)} words",
                  author_style),
        PageBreak(),
    ]
    for ch in chapters:
        story.append(Paragraph(ch.title, h1))
        story.append(Spacer(1, 12))
        for para in (ch.content or "").split("\n\n"):
            if para.strip():
                safe = para.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(safe, body))
                story.append(Spacer(1, 6))
        story.append(PageBreak())
    doc.build(story)
    return buf.getvalue()


def _manuscript_to_epub(title: str, author: str, chapters: list[Chapter]) -> bytes:
    from ebooklib import epub
    book = epub.EpubBook()
    book.set_identifier("asm-manuscript")
    book.set_title(title)
    book.set_language("en")
    book.add_author(author)
    # Cover / TOC
    book.toc = []
    chapters_epub = []
    for idx, ch in enumerate(chapters, start=1):
        c = epub.EpubHtml(title=ch.title, file_name=f"ch{idx}.xhtml")
        paras = "".join(f"<p>{p.strip()}</p>" for p in (ch.content or "").split("\n\n") if p.strip())
        c.set_content(f"<h1>{ch.title}</h1>{paras}")
        book.add_item(c)
        chapters_epub.append(c)
        book.toc.append(c)
    # Default spine
    book.spine = ["nav"] + chapters_epub
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    buf = io.BytesIO()
    epub.write_epub(buf, book, {})
    return buf.getvalue()


def _character_sheet_pdf(data: dict) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                     Table, TableStyle)
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter,
                            leftMargin=0.75*inch, rightMargin=0.75*inch,
                            topMargin=0.75*inch, bottomMargin=0.75*inch)
    styles = getSampleStyleSheet()
    name_style = ParagraphStyle("CharName", parent=styles["Title"],
                                alignment=TA_CENTER, fontSize=22, spaceAfter=6,
                                textColor=colors.HexColor(data.get("avatar_color") or "#6366f1"))
    role_style = ParagraphStyle("CharRole", parent=styles["Normal"],
                                alignment=TA_CENTER, fontSize=12, spaceAfter=18,
                                textColor=colors.grey)
    h2 = ParagraphStyle("CharH2", parent=styles["Heading2"],
                        fontSize=14, spaceBefore=12, spaceAfter=6,
                        textColor=colors.HexColor("#6366f1"))
    body = ParagraphStyle("CharBody", parent=styles["BodyText"],
                          fontSize=10, leading=14, spaceAfter=8)
    story = [
        Paragraph(data["name"], name_style),
        Paragraph(f"{data.get('role', '')} • {data.get('age', '')} • {data.get('gender', '')}",
                  role_style),
    ]
    # Aliases
    aliases = data.get("aliases") or []
    if aliases:
        story.append(Paragraph(f"<b>Aliases:</b> {', '.join(aliases)}", body))
    # Sections
    sections = [
        ("Physical", "physical"),
        ("Psychology", "psychology"),
        ("Background", "background"),
        ("Philosophy", "philosophy"),
        ("Story Role", "story_role"),
        ("Voice & Dialogue", "voice"),
        ("Notes", "notes"),
    ]
    for label, key in sections:
        v = data.get(key)
        if v:
            story.append(Paragraph(label, h2))
            safe = v.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(safe, body))
    # Quotes
    quotes = data.get("philosophy_quotes") or []
    if quotes:
        story.append(Paragraph("Notable Quotes", h2))
        for q in quotes:
            story.append(Paragraph(f'<i>"{q}"</i>', body))
    doc.build(story)
    return buf.getvalue()


def _character_sheet_docx(data: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    h = doc.add_heading(data["name"], 0)
    h.alignment = 1
    p = doc.add_paragraph()
    p.alignment = 1
    r = p.add_run(f"{data.get('role', '')} • {data.get('age', '')} • {data.get('gender', '')}")
    r.italic = True
    aliases = data.get("aliases") or []
    if aliases:
        doc.add_paragraph(f"Aliases: {', '.join(aliases)}")
    sections = [
        ("Physical", "physical"),
        ("Psychology", "psychology"),
        ("Background", "background"),
        ("Philosophy", "philosophy"),
        ("Story Role", "story_role"),
        ("Voice & Dialogue", "voice"),
        ("Notes", "notes"),
    ]
    for label, key in sections:
        v = data.get(key)
        if v:
            doc.add_heading(label, level=2)
            doc.add_paragraph(v)
    quotes = data.get("philosophy_quotes") or []
    if quotes:
        doc.add_heading("Notable Quotes", level=2)
        for q in quotes:
            doc.add_paragraph(f'"{q}"', style="Intense Quote")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _world_bible_pdf(title: str, entries: list[WorldEntry]) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, PageBreak)
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter,
                            leftMargin=1*inch, rightMargin=1*inch,
                            topMargin=1*inch, bottomMargin=1*inch)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("WBTitle", parent=styles["Title"],
                                 alignment=TA_CENTER, fontSize=24, spaceAfter=18)
    h1 = ParagraphStyle("WBH1", parent=styles["Heading1"],
                        fontSize=16, spaceBefore=12, spaceAfter=8,
                        textColor=colors.HexColor("#6366f1"))
    h2 = ParagraphStyle("WBH2", parent=styles["Heading2"],
                        fontSize=13, spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("WBBody", parent=styles["BodyText"],
                          fontSize=10, leading=14, spaceAfter=8)
    story = [Paragraph(f"{title} — World Bible", title_style), Spacer(1, 12)]

    by_type: dict[str, list] = {}
    for e in entries:
        by_type.setdefault(e.type, []).append(e)

    for t, group in by_type.items():
        story.append(Paragraph(t.upper().replace("_", " "), h1))
        for e in group:
            story.append(Paragraph(e.name, h2))
            if e.description:
                story.append(Paragraph(f"<i>{e.description}</i>", body))
            if e.content:
                safe = e.content.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(safe, body))
        story.append(PageBreak())
    doc.build(story)
    return buf.getvalue()


def _world_bible_docx(title: str, entries: list[WorldEntry]) -> bytes:
    from docx import Document
    doc = Document()
    h = doc.add_heading(f"{title} — World Bible", 0)
    h.alignment = 1
    by_type: dict[str, list] = {}
    for e in entries:
        by_type.setdefault(e.type, []).append(e)
    for t, group in by_type.items():
        doc.add_heading(t.upper().replace("_", " "), level=1)
        for e in group:
            doc.add_heading(e.name, level=2)
            if e.description:
                doc.add_paragraph(e.description)
            if e.content:
                doc.add_paragraph(e.content)
        doc.add_page_break()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_full_project(fmt: str = "docx",
                        include_secrets: bool = False) -> tuple[bytes, str, str]:
    """Export the entire project (manuscript + characters + world)."""
    # For brevity: produce a single combined DOCX/PDF.
    with read_session() as s:
        pid = current_project_id(s)
        from models.settings import Setting
        title = Setting.get(s, "story_title", "My Story")
        author = Setting.get(s, "story_author", "Author")
        chapters = list(s.scalars(
            select(Chapter).where(Chapter.project_id == pid)
            .order_by(Chapter.sort_order.asc())
        ).all())
        characters = list(s.scalars(
            select(Character).where(Character.project_id == pid)
            .order_by(Character.name.asc())
        ).all())
        world_entries = list(s.scalars(
            select(WorldEntry).where(WorldEntry.project_id == pid)
            .order_by(WorldEntry.type.asc(), WorldEntry.name.asc())
        ).all())
    if fmt == "docx":
        return _full_project_docx(title, author, chapters, characters, world_entries), \
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document", \
            f"{_safe_filename(title)}_full.docx"
    if fmt == "pdf":
        return _full_project_pdf(title, author, chapters, characters, world_entries), \
            "application/pdf", f"{_safe_filename(title)}_full.pdf"
    raise NotFoundError(f"Unknown format: {fmt}")


def _full_project_docx(title, author, chapters, characters, world_entries) -> bytes:
    from docx import Document
    doc = Document()
    h = doc.add_heading(title, 0)
    h.alignment = 1
    p = doc.add_paragraph()
    p.alignment = 1
    p.add_run(f"by {author}").italic = True
    doc.add_page_break()
    # TOC-ish
    doc.add_heading("Contents", level=1)
    doc.add_paragraph("Manuscript")
    doc.add_paragraph("Characters")
    doc.add_paragraph("World Bible")
    doc.add_page_break()
    # Manuscript
    doc.add_heading("Manuscript", level=1)
    for ch in chapters:
        doc.add_heading(ch.title, level=2)
        for para in (ch.content or "").split("\n\n"):
            if para.strip():
                doc.add_paragraph(para.strip())
        doc.add_page_break()
    # Characters
    doc.add_heading("Characters", level=1)
    for c in characters:
        doc.add_heading(c.name, level=2)
        doc.add_paragraph(f"Role: {c.role} • Age: {c.age} • Gender: {c.gender}")
        if c.physical:
            doc.add_paragraph(c.physical)
        if c.psychology:
            doc.add_paragraph(c.psychology)
    # World
    doc.add_heading("World Bible", level=1)
    by_type: dict[str, list] = {}
    for e in world_entries:
        by_type.setdefault(e.type, []).append(e)
    for t, group in by_type.items():
        doc.add_heading(t.upper().replace("_", " "), level=2)
        for e in group:
            doc.add_heading(e.name, level=3)
            if e.description:
                doc.add_paragraph(e.description)
            if e.content:
                doc.add_paragraph(e.content)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _full_project_pdf(title, author, chapters, characters, world_entries) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, PageBreak)
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter,
                            leftMargin=1*inch, rightMargin=1*inch,
                            topMargin=1*inch, bottomMargin=1*inch)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title", parent=styles["Title"],
                                 alignment=TA_CENTER, fontSize=24, spaceAfter=18)
    h1 = ParagraphStyle("H1", parent=styles["Heading1"],
                        fontSize=18, spaceBefore=12, spaceAfter=8,
                        textColor=colors.HexColor("#6366f1"))
    h2 = ParagraphStyle("H2", parent=styles["Heading2"],
                        fontSize=14, spaceBefore=10, spaceAfter=6)
    body = ParagraphStyle("Body", parent=styles["BodyText"],
                          fontSize=10, leading=14, spaceAfter=8)
    story = [
        Spacer(1, 2*inch),
        Paragraph(title, title_style),
        Paragraph(f"by {author}", body),
        PageBreak(),
    ]
    story.append(Paragraph("Manuscript", h1))
    for ch in chapters:
        story.append(Paragraph(ch.title, h2))
        for para in (ch.content or "").split("\n\n"):
            if para.strip():
                safe = para.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(safe, body))
    story.append(PageBreak())
    story.append(Paragraph("Characters", h1))
    for c in characters:
        story.append(Paragraph(c.name, h2))
        story.append(Paragraph(f"Role: {c.role} • Age: {c.age} • Gender: {c.gender}", body))
        if c.physical:
            story.append(Paragraph(c.physical.replace("&", "&amp;"), body))
        if c.psychology:
            story.append(Paragraph(c.psychology.replace("&", "&amp;"), body))
    story.append(PageBreak())
    story.append(Paragraph("World Bible", h1))
    by_type: dict[str, list] = {}
    for e in world_entries:
        by_type.setdefault(e.type, []).append(e)
    for t, group in by_type.items():
        story.append(Paragraph(t.upper().replace("_", " "), h2))
        for e in group:
            story.append(Paragraph(f"<b>{e.name}</b>", body))
            if e.description:
                story.append(Paragraph(f"<i>{e.description}</i>", body))
            if e.content:
                safe = e.content.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(safe, body))
    doc.build(story)
    return buf.getvalue()


# ---- Async export with progress ----

def async_export(entity_type: str, fmt: str, *, entity_id: str | None = None) -> str:
    """Start a background export task. Returns task_id.

    entity_type: "manuscript" | "world_bible" | "full_project" | "chapter" | "character"
    fmt: txt/md/html/docx/pdf/epub/json
    entity_id: required for chapter/character, ignored for manuscript/world_bible/full_project
    """
    task_id = uuid.uuid4().hex[:12]
    cache.set(f"export_progress:{task_id}", {
        "status": "processing",
        "progress": 0,
        "stage": "Starting...",
        "entity_type": entity_type,
        "fmt": fmt,
    }, ttl_seconds=600)

    thread = threading.Thread(
        target=_run_export_task,
        args=(task_id, entity_type, fmt, entity_id),
        daemon=True,
    )
    thread.start()
    return task_id


def _run_export_task(task_id: str, entity_type: str, fmt: str,
                     entity_id: str | None) -> None:
    """Run export in background thread, updating progress in cache."""
    try:
        _set_progress(task_id, 10, "Preparing export...")

        if entity_type == "manuscript":
            _set_progress(task_id, 30, "Gathering chapters...")
            content, mime, filename = export_manuscript(fmt)
        elif entity_type == "world_bible":
            _set_progress(task_id, 30, "Gathering world entries...")
            content, mime, filename = export_world_bible(fmt)
        elif entity_type == "full_project":
            _set_progress(task_id, 20, "Gathering all data...")
            content, mime, filename = export_full_project(fmt)
        elif entity_type == "chapter" and entity_id:
            _set_progress(task_id, 30, "Exporting chapter...")
            content, mime, filename = export_chapter(entity_id, fmt)
        elif entity_type == "character" and entity_id:
            _set_progress(task_id, 30, "Generating character sheet...")
            content, mime, filename = export_character_sheet(entity_id, fmt)
        else:
            raise NotFoundError(f"Unknown export type: {entity_type}")

        _set_progress(task_id, 80, "Finalizing file...")

        # Save to file
        export_dir = Config.EXPORT_DIR / task_id
        export_dir.mkdir(parents=True, exist_ok=True)
        file_path = export_dir / filename
        file_path.write_bytes(content)

        _set_progress(task_id, 100, "Done", status="done",
                      filename=filename, mime=mime, size=len(content))
        log.info("Export task %s completed: %s (%d bytes)", task_id, filename, len(content))

    except Exception as exc:
        log.exception("Export task %s failed", task_id)
        _set_progress(task_id, 0, f"Error: {exc}", status="error")


def _set_progress(task_id: str, progress: int, stage: str,
                  status: str = "processing", **extra: Any) -> None:
    """Update export progress in cache."""
    data = {
        "status": status,
        "progress": progress,
        "stage": stage,
    }
    data.update(extra)
    cache.set(f"export_progress:{task_id}", data, ttl_seconds=600)


def get_export_status(task_id: str) -> dict[str, Any] | None:
    """Get export task status. Returns None if task not found."""
    return cache.get(f"export_progress:{task_id}")


def get_export_file(task_id: str) -> tuple[bytes, str, str] | None:
    """Get the completed export file. Returns (content, mime, filename) or None."""
    status = get_export_status(task_id)
    if not status or status.get("status") != "done":
        return None
    filename = status.get("filename", "")
    mime = status.get("mime", "application/octet-stream")
    file_path = Config.EXPORT_DIR / task_id / filename
    if not file_path.exists():
        return None
    return file_path.read_bytes(), mime, filename
